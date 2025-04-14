#python -m unittest discover -s tests

import unittest
from pptx import Presentation
from orbat_gen.classes import Unit
from orbat_gen.powerpoint_gen import insert_app6, insert_unit_summary, insert_unit_eqp

class TestPowerPointGen(unittest.TestCase):

    def setUp(self):
        self.prs = Presentation()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.unit = Unit("Test Unit", "X")
        self.unit.equipments = {
            DummyEqpType("Rifle", "infantry"): 5,
            DummyEqpType("BTR", "infantry"): 7,
            DummyEqpType("Radio", "communication"): 1
        }

    def test_insert_app6_adds_shapes(self):
        nb_shapes_before = len(self.slide.shapes)
        insert_app6(self.slide, self.unit, 50,60)
        nb_shapes_after = len(self.slide.shapes)
        self.assertGreater(nb_shapes_after, nb_shapes_before)

    def test_insert_unit_summary_adds_table(self):
        insert_unit_summary(self.slide, self.unit, categories=['infantry'])
        table_shapes = [shape for shape in self.slide.shapes if shape.has_table]
        self.assertGreater(len(table_shapes), 0)

    def test_insert_unit_eqp_adds_table(self):
        insert_unit_eqp(self.slide, self.unit)
        table_shapes = [shape for shape in self.slide.shapes if shape.has_table]
        self.assertGreater(len(table_shapes), 0)


class DummyEqpType:
    def __init__(self, name, cat):
        self.name = name
        self.category = cat

    def __hash__(self):
        return hash(self.name)

    def __eq__(self, other):
        return isinstance(other, DummyEqpType) and self.name == other.name


if __name__ == '__main__':
    unittest.main()